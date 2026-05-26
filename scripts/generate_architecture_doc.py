from datetime import datetime
from docx import Document


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullets(doc: Document, items) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def main() -> None:
    doc = Document()

    add_heading(doc, "LLM Pipeline End-to-End Architecture Document", 0)
    doc.add_paragraph(
        "Prepared for: Project leadership and engineering stakeholders\n"
        f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )

    add_heading(doc, "1. Executive Summary", 1)
    doc.add_paragraph(
        "This project implements an end-to-end pipeline that creates a training dataset from source code, "
        "collects human labels in Label Studio, prepares model-ready training data, and trains a GraphCodeBERT-based model."
    )

    add_heading(doc, "2. High-Level Architecture", 1)
    doc.add_paragraph("Label Studio (Human labels)")
    doc.add_paragraph("↓")
    doc.add_paragraph("Export labeled dataset (JSON)")
    doc.add_paragraph("↓")
    doc.add_paragraph("Convert to training format (query, code, label)")
    doc.add_paragraph("↓")
    doc.add_paragraph("Train model (GraphCodeBERT / embedding model)")
    doc.add_paragraph("↓")
    doc.add_paragraph("Evaluate model")
    doc.add_paragraph("↓")
    doc.add_paragraph("Deploy / integrate into search system")

    add_heading(doc, "3. Current Implemented Flow", 1)
    add_numbered(
        doc,
        [
            "Step 1 (`pipeline.py`): Repo sync, function extraction, cleaning, dataset build, DVC tracking, upload to Label Studio.",
            "Step 2: Human annotation in Label Studio.",
            "Step 3 (`training_pipeline.py`): Export labels, convert labels, auto-label, merge, and train GraphCodeBERT.",
            "Master execution (`run_all.py`): one command to run full flow or skip Step 1.",
        ],
    )

    add_heading(doc, "4. Component-Level Design", 1)
    add_heading(doc, "4.1 Data Preparation Pipeline", 2)
    add_bullets(
        doc,
        [
            "`scripts/repo_sync.py`: clones/pulls source repository into `data/raw/repo`.",
            "`scripts/extract_functions.py`: scans Python files and extracts function signatures.",
            "`scripts/clean_data.py`: creates enriched records (`hard_query`, `soft_query`, `code`, `instruction`).",
            "`scripts/build_dataset.py`: saves Hugging Face dataset to `data/processed/instruct_v1`.",
            "DVC tracking: versions dataset via `python -m dvc add data/processed/instruct_v1`.",
            "`scripts/upload_labelstudio.py`: logs in, recreates project, uploads tasks, stores `project_id.txt`.",
        ],
    )

    add_heading(doc, "4.2 Labeling and Training Pipeline", 2)
    add_bullets(
        doc,
        [
            "`scripts/export_labels.py`: exports Label Studio project labels to `labelstudio_export.json`.",
            "`scripts/convert_labels.py`: converts export format into `train.json` with (`query`, `code`, `label`).",
            "`scripts/auto_label.py`: augments with weak labels (`auto.json`); default fast heuristic with optional MiniLM.",
            "`scripts/merge.py`: merges human and auto labels into `final_train.json`.",
            "`scripts/train.py`: trains GraphCodeBERT on merged dataset with configurable limits.",
        ],
    )

    add_heading(doc, "5. Data Contracts and Artifacts", 1)
    add_bullets(
        doc,
        [
            "`data/processed/cleaned.json`: cleaned candidate data before labeling.",
            "`labelstudio_export.json`: raw Label Studio export.",
            "`train.json`: human-labeled training rows.",
            "`auto.json`: auto-labeled rows.",
            "`final_train.json`: merged training dataset.",
            "`project_id.txt`: current Label Studio project id.",
        ],
    )

    add_heading(doc, "6. Execution Modes", 1)
    doc.add_paragraph("Recommended team commands:")
    add_bullets(
        doc,
        [
            "Full run: `.\\.venv\\Scripts\\python.exe run_all.py --mode fast`",
            "Retrain only: `.\\.venv\\Scripts\\python.exe run_all.py --mode fast --skip-step1`",
            "Higher-quality mode: `.\\.venv\\Scripts\\python.exe run_all.py --mode full`",
        ],
    )
    doc.add_paragraph("Fast mode defaults:")
    add_bullets(
        doc,
        [
            "MAX_AUTO_SAMPLES=80",
            "MAX_TRAIN_SAMPLES=24",
            "TRAIN_EPOCHS=1",
            "TRAIN_BATCH_SIZE=8",
            "USE_MINILM=0",
        ],
    )

    add_heading(doc, "7. Tools and Infrastructure", 1)
    add_bullets(
        doc,
        [
            "Python + virtual environment",
            "Label Studio (`http://localhost:8080`)",
            "DVC for data versioning",
            "Hugging Face model downloads (token optional but recommended)",
            "GitHub repository for collaboration and version control",
        ],
    )

    add_heading(doc, "8. What Is Already Achieved", 1)
    add_bullets(
        doc,
        [
            "End-to-end MVP flow is operational.",
            "One-command execution is available via `run_all.py`.",
            "Human labels are successfully exported and included in model training.",
            "Training pipeline completes in fast mode on local machine.",
        ],
    )

    add_heading(doc, "9. Gaps to Reach Production Readiness", 1)
    add_bullets(
        doc,
        [
            "Evaluation stage needs to be formalized (metrics, validation split, acceptance thresholds).",
            "Model artifact saving/versioning should be added after training.",
            "Deployment/integration into search system needs implementation design and rollout plan.",
            "Secrets management should move to secure `.env`/secret manager usage only.",
            "CI automation via GitHub Actions is recommended for lint/test/validation and retraining triggers.",
        ],
    )

    add_heading(doc, "10. Proposed Next Milestones", 1)
    add_numbered(
        doc,
        [
            "Add `evaluate.py` with precision/recall/F1 and baseline comparison.",
            "Save trained model checkpoints and metadata with version tags.",
            "Implement inference service or batch scoring endpoint for search integration.",
            "Add GitHub Actions workflow for CI and optional scheduled pipeline run.",
            "Publish team runbook and on-call troubleshooting guide.",
        ],
    )

    output = "LLM_Pipeline_Architecture_Deep_Dive.docx"
    doc.save(output)
    print(f"Created {output}")


if __name__ == "__main__":
    main()
